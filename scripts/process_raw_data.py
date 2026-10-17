#!/usr/bin/env python3
"""附表五（信用卡重要業務及財務資訊揭露）原始資料清理與標準化程式。

依據：
- CLAUDE.md
- docs/data-dictionary.md
- docs/calculation-rules.md
- docs/raw-data-inspection.md

本程式只讀取 data/raw 內的 .xlsx 原始檔，不修改、不覆寫任何原始檔，
將資料標準化後輸出至 data/processed。任一關鍵資料品質驗證失敗時，
不輸出任何看似成功的正式資料，並以非 0 狀態結束。
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RAW_DIR = PROJECT_ROOT / "data" / "raw"
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"

# ---------------------------------------------------------------------------
# 正式規格（依 docs/data-dictionary.md，原始中文欄位名稱逐字保留）
# ---------------------------------------------------------------------------

COLUMN_SPEC = [
    ("A", "金融機構名稱", "institution_name"),
    ("B", "流通卡數　　　", "cards_in_force"),
    ("C", "有效卡數　　　", "active_cards"),
    ("D", "當月發卡數　　　", "cards_issued_month"),
    ("E", "當月停卡數　　　", "cards_cancelled_month"),
    ("F", "循環信用餘額　　", "revolving_credit_balance"),
    ("G", "未到期分期付款餘額　", "installment_balance_not_due"),
    ("H", "當月簽帳金額　　", "purchase_amount_month"),
    ("I", "當月預借現金金額　　", "cash_advance_amount_month"),
    ("J", "逾期三個月以上帳款占應收帳款餘額(含催收款)之比率(%)", "delinquency_ratio_3m"),
    ("K", "逾期六個月以上帳款占應收帳款餘額(含催收款)之比率(%)", "delinquency_ratio_6m"),
    ("L", "備抵呆帳提足率(%) ", "reserve_coverage_ratio"),
    ("M", "當月轉銷呆帳金額　 ", "charge_off_amount_month"),
    ("N", "當年度轉銷呆帳金額累計至資料月份", "charge_off_amount_ytd"),
]

# 可加總欄位（比率欄位 J、K、L 不可加總，依 docs/calculation-rules.md 第 4 項）
SUMMABLE_FIELDS = [
    "cards_in_force", "active_cards", "cards_issued_month", "cards_cancelled_month",
    "revolving_credit_balance", "installment_balance_not_due", "purchase_amount_month",
    "cash_advance_amount_month", "charge_off_amount_month", "charge_off_amount_ytd",
]

MARKET_TOTAL_NAME = "總計"
FAR_EASTERN_BANK_NAME = "遠東國際商業銀行"

# 正式機構名單（依 docs/raw-data-inspection.md 第五、六階段唯讀檢查結果）
BANK_NAMES = [
    "臺灣銀行", "臺灣土地銀行", "合作金庫商業銀行", "第一商業銀行", "華南商業銀行",
    "彰化商業銀行", "上海商業儲蓄銀行", "台北富邦商業銀行", "國泰世華商業銀行", "高雄銀行",
    "兆豐國際商業銀行", "花旗(台灣)商業銀行", "臺灣中小企業銀行", "渣打國際商業銀行", "台中商業銀行",
    "滙豐(台灣)商業銀行", "華泰商業銀行", "臺灣新光商業銀行", "陽信商業銀行", "三信商業銀行",
    "聯邦商業銀行", "遠東國際商業銀行", "元大商業銀行", "永豐商業銀行", "玉山商業銀行",
    "凱基商業銀行", "星展(台灣)商業銀行", "台新國際商業銀行", "安泰商業銀行", "中國信託商業銀行",
]
NON_BANK_ISSUER_NAMES = ["台灣樂天信用卡股份有限公司", "台灣美國運通國際(股)公司"]

ISSUER_TYPE = {name: "bank" for name in BANK_NAMES}
ISSUER_TYPE.update({name: "non_bank_issuer" for name in NON_BANK_ISSUER_NAMES})

EXPECTED_ISSUER_COUNT = 32
EXPECTED_BANK_COUNT = 30
EXPECTED_NON_BANK_COUNT = 2

FIELDNAMES = [
    "period", "roc_year", "month", "issuer_name_raw", "issuer_name", "issuer_type",
    "institution_name",
    "cards_in_force", "active_cards", "cards_issued_month", "cards_cancelled_month",
    "revolving_credit_balance", "installment_balance_not_due", "purchase_amount_month",
    "cash_advance_amount_month", "delinquency_ratio_3m", "delinquency_ratio_6m",
    "reserve_coverage_ratio", "charge_off_amount_month", "charge_off_amount_ytd",
    "source_file", "source_sha256",
]

METADATA_FIELDNAMES = [
    "period", "source_file", "source_sha256", "sheet_name", "source_unit",
    "issuer_count", "bank_count", "non_bank_issuer_count", "source_url", "validation_status",
    "duplicate_sources", "all_source_files", "selected_source_file", "dedup_reason",
    "correction_markers",
]

# 金管會原始檔正式修正註記（依原始檔備註：「r表示於申報日後，曾異動該項資料。」）
CORRECTION_MARKER_MEANING = "原始檔註明該項資料於申報日後曾異動"

MONTH_PATTERN = re.compile(r"(\d{2,3})[\s　]*年[\s　]*(\d{1,2})[\s　]*月")
# 檔名開頭 5 碼民國年月（民國年 3 碼＋月份 2 碼，限 01～12），底線可有可無，
# 但後面不得緊接其他數字（避免誤判更長的數字序列），例如：
#   11401_信用卡重要資訊揭露.xlsx（有底線）
#   11401信用卡重要資訊揭露.xlsx（無底線，原始檔既有命名，不得要求改名）
FILENAME_PREFIX_PATTERN = re.compile(r"^(\d{3})(0[1-9]|1[0-2])_?(?!\d)")
# 「5 碼民國年月＋底線」為優先命名格式，用於同月重複來源的去重選用
PREFERRED_FILENAME_PATTERN = re.compile(r"^(\d{3})(0[1-9]|1[0-2])_")


class ValidationError(Exception):
    """資料品質驗證失敗。發生時不得輸出看似成功的正式資料。"""


# ---------------------------------------------------------------------------
# 工具函式
# ---------------------------------------------------------------------------

def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_roc_year_month_from_text(text: str) -> Optional[tuple[int, int]]:
    if not text:
        return None
    m = MONTH_PATTERN.search(text)
    if not m:
        return None
    month = int(m.group(2))
    if not 1 <= month <= 12:
        return None
    return int(m.group(1)), month


def parse_roc_year_month_from_filename(filename: str) -> Optional[tuple[int, int]]:
    """從檔名辨識民國年月，同時支援：
    - 「ＸＸＸ年ＸＸ月」文字型檔名（如 附表5_115年7月信用卡重要資訊揭露.xlsx）
    - 「5 碼民國年月＋底線」檔名（如 11401_信用卡重要資訊揭露.xlsx）
    - 「5 碼民國年月，無底線」檔名（如 11401信用卡重要資訊揭露.xlsx，原始檔既有命名）
    無法辨識時回傳 None，由呼叫端清楚報錯並指出檔名。"""
    m = MONTH_PATTERN.search(filename)
    if m:
        month = int(m.group(2))
        if 1 <= month <= 12:
            return int(m.group(1)), month
    m = FILENAME_PREFIX_PATTERN.match(filename)
    if m:
        return int(m.group(1)), int(m.group(2))
    return None


def is_preferred_filename(filename: str) -> bool:
    """「5 碼民國年月＋底線」為優先命名格式，供同月重複來源去重時選用。"""
    return PREFERRED_FILENAME_PATTERN.match(filename) is not None


def roc_to_period(roc_year: int, month: int) -> str:
    western_year = roc_year + 1911
    return f"{western_year:04d}-{month:02d}"


def discover_raw_files() -> list[Path]:
    """只回傳 data/raw 內副檔名為 .xlsx 的附表五檔案，忽略 README、ZIP、ODS
    及 Excel 暫存檔（以 ~$ 開頭）。"""
    files = []
    for p in sorted(RAW_DIR.glob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        files.append(p)
    return files


def _row_values_by_col(row) -> dict[int, object]:
    """將 iter_rows 回傳的一列儲存格轉為 {欄位編號: 值} 字典。
    read_only 模式下超出資料範圍的儲存格為 EmptyCell（無 .row/.column 屬性，
    value 恆為 None），故一律以列內位置（1 起算）作為欄位編號，不依賴
    cell.column 屬性。"""
    return {idx: cell.value for idx, cell in enumerate(row, start=1)}


def find_header_row(ws) -> int:
    for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=10), start=1):
        for cell in row:
            if cell.value == "金融機構名稱":
                return row_num
    raise ValidationError("找不到表頭列（欄位「金融機構名稱」），無法定位欄位")


def build_column_index(ws, header_row: int) -> dict[str, int]:
    """依原始中文欄位名稱（逐字比對）建立英文欄位名稱到欄位編號的對照，
    不依賴欄位固定位置。"""
    text_to_field = {orig: eng for _, orig, eng in COLUMN_SPEC}
    col_idx: dict[str, int] = {}
    header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
    for idx, cell in enumerate(header_cells, start=1):
        if isinstance(cell.value, str) and cell.value in text_to_field:
            col_idx[text_to_field[cell.value]] = idx
    missing = [eng for _, _, eng in COLUMN_SPEC if eng not in col_idx]
    if missing:
        raise ValidationError(f"表頭欄位比對失敗，缺少以下標準欄位（依原始中文名稱逐字比對）：{missing}")
    return col_idx


def find_text_matching(ws, max_row: int, predicate) -> Optional[str]:
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for cell in row:
            if isinstance(cell.value, str) and predicate(cell.value):
                return cell.value
    return None


# 數字文字格式：開頭至多 1 個金管會正式修正註記字母 r/R（可接 0 個以上空白），
# 其後為純數字（可含千分位逗號、至多 1 個小數點）。不得移除其他位置的英文字母，
# 例如 "x 9,001,262"、"r abc"、"9r001262" 皆不符合此格式。
NUMERIC_TEXT_PATTERN = re.compile(r"^([rR])?\s*([\d,]+(?:\.\d+)?)$")


def parse_numeric_cell_text(text: str) -> Optional[tuple[float, Optional[str]]]:
    """嘗試將儲存格文字解析為數字，支援金管會正式修正註記（r/R）與千分位逗號。

    成功時回傳 (parsed_value, marker)：marker 為 'r'／'R'（偵測到修正註記）或
    None（純數字文字，無註記）。無法解析為合法數字時回傳 None，由呼叫端決定
    如何報錯（不得在此處猜測或靜默忽略無法解析的內容）。"""
    m = NUMERIC_TEXT_PATTERN.match(text.strip())
    if not m:
        return None
    marker, number_text = m.group(1), m.group(2)
    cleaned = number_text.replace(",", "")
    if not re.fullmatch(r"\d+(\.\d+)?", cleaned):
        return None
    parsed = float(cleaned) if "." in cleaned else int(cleaned)
    return parsed, marker


def extract_fields(
    cells_by_col: dict[int, object], col_idx: dict[str, int], filename: str
) -> tuple[dict, list[dict]]:
    """解析 B～N 欄數值，回傳 (欄位值字典, 修正註記清單)。

    修正註記清單中每筆為 {"field", "raw_value", "parsed_value", "marker"}，
    由呼叫端補上機構／官方合計列等脈絡後，彙整寫入 source_metadata.json，
    不得只將 r/R 註記靜默移除而不留痕跡。"""
    result: dict = {}
    markers: list[dict] = []
    for _, _, eng in COLUMN_SPEC:
        if eng == "institution_name":
            continue
        value = cells_by_col.get(col_idx[eng])
        if value is None:
            result[eng] = None
            continue
        if isinstance(value, (int, float)):
            result[eng] = value
            continue
        if isinstance(value, str):
            parsed = parse_numeric_cell_text(value)
            if parsed is not None:
                parsed_value, marker = parsed
                result[eng] = parsed_value
                if marker is not None:
                    markers.append({
                        "field": eng,
                        "raw_value": value,
                        "parsed_value": parsed_value,
                        "marker": marker,
                    })
                continue
        raise ValidationError(f"{filename}：欄位 {eng} 出現無法解析為數字的文字內容：{value!r}")
    return result, markers


# ---------------------------------------------------------------------------
# 單一檔案處理
# ---------------------------------------------------------------------------

def process_file(path: Path) -> tuple[list[dict], dict, dict]:
    sha_before = sha256_of(path)

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        sheet_name = ws.title
        header_row = find_header_row(ws)
        col_idx = build_column_index(ws, header_row)

        unit_text = find_text_matching(ws, header_row, lambda s: "單位" in s)
        period_text = find_text_matching(ws, header_row, lambda s: MONTH_PATTERN.search(s) is not None)

        sheet_ym = parse_roc_year_month_from_text(period_text) if period_text else None
        if sheet_ym is None:
            raise ValidationError(f"{path.name}：無法從工作表內容辨識資料年月（找不到「ＸＸＸ年ＸＸ月」文字）")

        filename_ym = parse_roc_year_month_from_filename(path.name)
        if filename_ym is None:
            raise ValidationError(f"{path.name}：無法從檔名辨識資料年月，供交叉驗證")

        if sheet_ym != filename_ym:
            raise ValidationError(
                f"{path.name}：工作表年月（民國{sheet_ym[0]}年{sheet_ym[1]}月）"
                f"與檔名年月（民國{filename_ym[0]}年{filename_ym[1]}月）不一致，停止處理"
            )

        roc_year, month = sheet_ym
        period = roc_to_period(roc_year, month)

        institutions: list[dict] = []
        market_total: Optional[dict] = None
        seen_names: set[str] = set()
        correction_markers: list[dict] = []

        for row in ws.iter_rows(min_row=header_row + 1):
            cells_by_col = _row_values_by_col(row)
            name_val = cells_by_col.get(col_idx["institution_name"])
            if name_val is None:
                continue
            name_str = str(name_val).strip()
            if not name_str:
                continue

            if name_str == MARKET_TOTAL_NAME:
                if market_total is not None:
                    raise ValidationError(f"{path.name}：官方合計列（總計）出現超過一次")
                fields, field_markers = extract_fields(cells_by_col, col_idx, path.name)
                market_total = {
                    "issuer_name_raw": str(name_val),
                    "issuer_name": name_str,
                    "issuer_type": "market_total",
                    "institution_name": name_str,
                    **fields,
                }
                for fm in field_markers:
                    correction_markers.append({
                        "month": period,
                        "source_file": path.name,
                        "row_type": "market_total",
                        "institution": MARKET_TOTAL_NAME,
                        "field": fm["field"],
                        "raw_value": fm["raw_value"],
                        "parsed_value": fm["parsed_value"],
                        "marker": fm["marker"],
                        "meaning": CORRECTION_MARKER_MEANING,
                    })
                continue

            if name_str in ISSUER_TYPE:
                if name_str in seen_names:
                    raise ValidationError(f"{path.name}：機構名稱重複出現：{name_str}")
                seen_names.add(name_str)
                fields, field_markers = extract_fields(cells_by_col, col_idx, path.name)
                institutions.append({
                    "issuer_name_raw": str(name_val),
                    "issuer_name": name_str,
                    "issuer_type": ISSUER_TYPE[name_str],
                    "institution_name": name_str,
                    **fields,
                })
                for fm in field_markers:
                    correction_markers.append({
                        "month": period,
                        "source_file": path.name,
                        "row_type": "issuer",
                        "institution": name_str,
                        "field": fm["field"],
                        "raw_value": fm["raw_value"],
                        "parsed_value": fm["parsed_value"],
                        "marker": fm["marker"],
                        "meaning": CORRECTION_MARKER_MEANING,
                    })
                continue

            # 其他非機構文字（備註列等），依規則不得當成機構資料，略過
            continue

        if len(institutions) != EXPECTED_ISSUER_COUNT:
            raise ValidationError(
                f"{path.name}：發卡機構筆數為 {len(institutions)}，預期 {EXPECTED_ISSUER_COUNT} 家"
            )

        bank_count = sum(1 for r in institutions if r["issuer_type"] == "bank")
        non_bank_count = sum(1 for r in institutions if r["issuer_type"] == "non_bank_issuer")
        if bank_count != EXPECTED_BANK_COUNT or non_bank_count != EXPECTED_NON_BANK_COUNT:
            raise ValidationError(
                f"{path.name}：銀行/非銀行發卡機構家數不符，銀行 {bank_count} 家、"
                f"非銀行 {non_bank_count} 家（預期 {EXPECTED_BANK_COUNT}／{EXPECTED_NON_BANK_COUNT}）"
            )

        far_eastern_count = sum(1 for r in institutions if r["issuer_name"] == FAR_EASTERN_BANK_NAME)
        if far_eastern_count != 1:
            raise ValidationError(
                f"{path.name}：「{FAR_EASTERN_BANK_NAME}」出現 {far_eastern_count} 次，預期恰好 1 次"
            )

        if market_total is None:
            raise ValidationError(f"{path.name}：找不到官方合計列（總計）")

        recon_errors: dict[str, dict] = {}
        for eng in SUMMABLE_FIELDS:
            values = [r[eng] for r in institutions]
            if any(v is None for v in values) or market_total[eng] is None:
                continue  # 含缺值時不比對加總，缺值不得補 0
            total = sum(values)
            if total != market_total[eng]:
                recon_errors[eng] = {
                    "sum_of_32": total,
                    "market_total": market_total[eng],
                    "diff": total - market_total[eng],
                }
        if recon_errors:
            raise ValidationError(
                f"{path.name}：32 家發卡機構明細加總與官方合計列不一致（資料異常，未自行修改）：{recon_errors}"
            )

        sha_after = sha256_of(path)
        if sha_after != sha_before:
            raise ValidationError(
                f"{path.name}：處理前後 SHA-256 不一致，原始檔可能已被修改"
                f"（處理前：{sha_before}，處理後：{sha_after}）"
            )

        for r in institutions:
            r.update(period=period, roc_year=roc_year, month=month,
                      source_file=path.name, source_sha256=sha_after)
        market_total.update(period=period, roc_year=roc_year, month=month,
                             source_file=path.name, source_sha256=sha_after)

        metadata = {
            "period": period,
            "source_file": path.name,
            "source_sha256": sha_after,
            "sheet_name": sheet_name,
            "source_unit": unit_text,
            "issuer_count": len(institutions),
            "bank_count": bank_count,
            "non_bank_issuer_count": non_bank_count,
            "source_url": None,
            "validation_status": "PASS",
            "correction_markers": correction_markers,
        }
        return institutions, market_total, metadata
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# 同月重複來源防護
# ---------------------------------------------------------------------------

def resolve_duplicate_group(period: str, group: list[tuple[list[dict], dict, dict]]) -> tuple[tuple[list[dict], dict, dict], dict]:
    """同一資料年月若對應多個來源檔案，依下列規則處理：
    - 僅 1 個來源：直接採用，並記錄「無重複」中繼資料。
    - 多個來源但 SHA-256 皆相同：視為同一份資料的不同命名／下載管道，
      只處理一次，優先採用「5 碼民國年月＋底線」命名的檔案。
    - 多個來源但 SHA-256 不同：資料內容確實不同，不得自行選擇，
      拋出 ValidationError 並列出月份、檔名與 SHA-256。
    """
    all_source_files = sorted(m["source_file"] for _, _, m in group)

    if len(group) == 1:
        chosen = group[0]
        dup_info = {
            "duplicate_sources": False,
            "all_source_files": all_source_files,
            "selected_source_file": chosen[2]["source_file"],
            "dedup_reason": None,
        }
        return chosen, dup_info

    hashes = {m["source_sha256"] for _, _, m in group}
    if len(hashes) > 1:
        detail = {m["source_file"]: m["source_sha256"] for _, _, m in group}
        raise ValidationError(
            f"{period}：同一資料月份存在內容不同的多個原始檔，無法自動選擇其中一份。"
            f"請人工確認後再處理。檔案與 SHA-256：{detail}"
        )

    # SHA-256 相同 -> 安全去重，優先選用「5 碼民國年月＋底線」命名的檔案
    preferred = [item for item in group if is_preferred_filename(item[2]["source_file"])]
    if preferred:
        chosen = sorted(preferred, key=lambda item: item[2]["source_file"])[0]
        reason = "SHA-256 相同，優先採用「5碼民國年月＋底線」命名檔案"
    else:
        chosen = sorted(group, key=lambda item: item[2]["source_file"])[0]
        reason = "SHA-256 相同，命名皆非「5碼民國年月＋底線」標準格式，依檔名排序選用"

    dup_info = {
        "duplicate_sources": True,
        "all_source_files": all_source_files,
        "selected_source_file": chosen[2]["source_file"],
        "dedup_reason": reason,
    }
    return chosen, dup_info


# ---------------------------------------------------------------------------
# 整體資料集建立與輸出
# ---------------------------------------------------------------------------

def build_dataset() -> tuple[list[dict], list[dict], list[dict]]:
    files = discover_raw_files()
    if not files:
        raise ValidationError(f"{RAW_DIR} 內找不到任何 .xlsx 原始檔")

    # 逐檔處理（含表頭、機構名單、加總等驗證），再依資料年月分組
    file_results: list[tuple[list[dict], dict, dict]] = []
    for path in files:
        file_results.append(process_file(path))

    groups: dict[str, list[tuple[list[dict], dict, dict]]] = {}
    for item in file_results:
        period = item[2]["period"]
        groups.setdefault(period, []).append(item)

    all_institutions: list[dict] = []
    all_market_totals: list[dict] = []
    all_metadata: list[dict] = []

    for period in sorted(groups.keys()):
        (institutions, market_total, metadata), dup_info = resolve_duplicate_group(period, groups[period])
        metadata = {**metadata, **dup_info}
        all_institutions.extend(institutions)
        all_market_totals.append(market_total)
        all_metadata.append(metadata)

    # 依 period 由舊至新排序；Python sort 為穩定排序，同一 period 內原始檔機構順序不變
    all_institutions.sort(key=lambda r: r["period"])
    all_market_totals.sort(key=lambda r: r["period"])
    all_metadata.sort(key=lambda r: r["period"])

    return all_institutions, all_market_totals, all_metadata


def _ordered(records: list[dict], fieldnames: list[str]) -> list[dict]:
    return [{k: r.get(k) for k in fieldnames} for r in records]


def write_csv(records: list[dict], path: Path, fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for r in records:
            row = {k: ("" if r.get(k) is None else r.get(k)) for k in fieldnames}
            writer.writerow(row)


def write_json(records: list[dict], path: Path) -> None:
    with path.open("w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
        f.write("\n")


def write_outputs(monthly: list[dict], market: list[dict], metadata: list[dict], out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)

    monthly_ordered = _ordered(monthly, FIELDNAMES)
    market_ordered = _ordered(market, FIELDNAMES)
    metadata_ordered = _ordered(metadata, METADATA_FIELDNAMES)

    write_csv(monthly_ordered, out_dir / "credit_card_monthly.csv", FIELDNAMES)
    write_json(monthly_ordered, out_dir / "credit_card_monthly.json")
    write_csv(market_ordered, out_dir / "market_totals.csv", FIELDNAMES)
    write_json(market_ordered, out_dir / "market_totals.json")
    write_json(metadata_ordered, out_dir / "source_metadata.json")


def main() -> int:
    try:
        monthly, market, metadata = build_dataset()
    except ValidationError as exc:
        print("[錯誤] 資料品質驗證失敗，未輸出任何正式資料。", file=sys.stderr)
        print(str(exc), file=sys.stderr)
        return 1

    write_outputs(monthly, market, metadata, PROCESSED_DIR)

    print(f"處理完成：發卡機構明細 {len(monthly)} 筆、市場合計 {len(market)} 筆，涵蓋 {len(metadata)} 個月份。")
    for m in metadata:
        print(
            f"  {m['period']}｜來源：{m['source_file']}｜"
            f"銀行 {m['bank_count']} 家、非銀行 {m['non_bank_issuer_count']} 家｜"
            f"驗證狀態：{m['validation_status']}"
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
